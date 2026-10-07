from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
out=root/'outputs'; out.mkdir(exist_ok=True)
d=Document(); s=d.sections[0]
s.top_margin=s.bottom_margin=Inches(.7); s.left_margin=s.right_margin=Inches(.8)
s.page_width=Inches(8.5); s.page_height=Inches(11)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
st=d.styles['Normal']; st.font.size=Pt(11); st.paragraph_format.space_after=Pt(7); st.paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(26)
d.styles['Heading 1'].font.size=Pt(17)
d.styles['Heading 2'].font.size=Pt(12)
for name in ['Heading 1','Heading 2']:
 d.styles[name].paragraph_format.space_before=Pt(12)
 d.styles[name].paragraph_format.space_after=Pt(7)
footer=s.footer.paragraphs[0]; footer.alignment=2
r=footer.add_run(); r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1, cols=len(headers)); t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 for i,txt in enumerate(headers): t.rows[0].cells[i].text=txt
 for row in rows:
  for c,txt in zip(t.add_row().cells,row): c.text=str(txt)
 for ri,row in enumerate(t.rows):
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]); c.vertical_alignment=1
   tcPr=c._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'DCE6EF' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF')); tcPr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
   tcPr.append(borders)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); mar.append(e)
   tcPr.append(mar)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(2); par.paragraph_format.line_spacing=1.0
    for r in par.runs: r.font.size=Pt(10); r.bold=ri==0
  trPr=row._tr.get_or_add_trPr(); ns=OxmlElement('w:cantSplit'); trPr.append(ns)
 rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
 d.add_paragraph().paragraph_format.space_after=Pt(0)

d.add_paragraph('Propuesta de evaluación de pitcheo para Diablos Rojos', 'Title')
d.add_paragraph('Hackathon 2026 · Plan de trabajo para cuatro integrantes', 'Subtitle')
p('27 de septiembre de 2026')
h('Planteamiento del proyecto')
p('Proponemos desarrollar un sistema que evalúe la calidad de cada lanzamiento y compare la adaptación del arsenal de un pitcher a distintos ambientes de altitud. La meta es apoyar una decisión concreta de scouting: entre varios candidatos, identificar qué repertorio parece más adecuado para el Harp Helú, qué características explican la recomendación y cuánta incertidumbre existe.')
p('La primera versión se construirá con las categorías de altitud disponibles en el dataset. La comparación específica del Harp Helú quedará condicionada a recibir un identificador de estadio o un mapeo oficial de partidos. No presentaremos una categoría de altitud como si identificara por sí sola a CDMX.')
h('Producto propuesto')
p('Un dashboard de evaluación y simulación permitirá seleccionar un pitcher, revisar su repertorio, consultar Stuff+, Location+ y Pitching+, y comparar escenarios de ambiente. Cada resultado mostrará las características que lo impulsan, su tamaño de muestra y sus límites de uso.')
h('Objetivo en dos semanas')
p('Entregar un prototipo reproducible con evaluación fuera de muestra, análisis de arsenal y ambiente, tres casos de scouting y una explicación técnica clara. La prioridad será demostrar utilidad con evidencia; no se promete un efecto causal de altitud ni una mejora predictiva antes de medirla.')
h('Entregables')
for x in ['Scores por lanzamiento, pitcher y arsenal, con una escala de referencia común centrada en 100.', 'Comparador de ambientes y mapa de movimiento con incertidumbre y advertencias de muestra insuficiente.', 'Perfil físico recomendado y tres casos contrastantes: adaptación favorable, desfavorable y evidencia insuficiente.', 'Documento técnico de 6 a 10 páginas, resultados de validación y demostración para la entrevista del jurado.']:
 d.add_paragraph(x,'List Bullet')

page('Datos disponibles y evidencia inicial')
p('La auditoría revisó toda la representación Parquet y el diccionario incluidos en el ZIP. El archivo también contiene versiones RDS y pickle, cuya equivalencia no se comprobó. Los siguientes resultados describen los datos; todavía no se ha entrenado un modelo.')
table(['Recurso','Cobertura'],[
('Lanzamientos y variables','635,002 lanzamientos y 84 columnas'),('Entidades identificadas','1,134 pitchers, 864 bateadores y 2,127 partidos'),('Temporadas','2024: 243,254 · 2025: 256,144 · 2026: 135,604 lanzamientos'),('Integridad básica','635,002 PitchUID únicos; sin duplicados en ese identificador'),('Variables físicas principales','99.76% de filas completas en las ocho variables revisadas'),('Comparación ambiental','865 pitchers presentes en más de una categoría de altitud')],[2.05,4.85])
p('Contamos con velocidad, spin, movimiento vertical y horizontal, release, extensión, ubicación, conteo, lado del bateador y resultados. Hay mediciones de contacto que permiten explorar contacto débil o fuerte; sus definiciones deberán acordarse antes de entrenar.')
h('Primer hallazgo que orienta la propuesta')
p('Al comparar el mismo pitcher, año y tipo de lanzamiento, con al menos 30 observaciones en cada categoría, aparecen diferencias de movimiento entre Extreme Altitude y No Altitude. Se calculó la media por grupo y ambiente y después la mediana de las diferencias, dando el mismo peso a cada grupo.')
table(['Lanzamiento','Grupos comparables','Diferencia mediana de IVB'],[('Recta de cuatro costuras','450','−3.84 pulgadas'),('Sinker','178','−2.46 pulgadas'),('Cambio','101','−2.19 pulgadas'),('Slider','271','−0.74 pulgadas'),('Curva','64','+2.02 pulgadas')],[2.65,1.65,2.6])
p('IVB significa movimiento vertical inducido. La diferencia es altitud extrema menos No Altitude. El signo positivo de la curva no implica mejor calidad. Estas asociaciones aún no controlan spin, velocidad, release, selección ni sensores; no son correcciones físicas ni deltas Stuff+ validados.')

page('Diseño del modelo y factores diferenciadores')
h('Tres capas con funciones distintas')
p('Stuff+ evaluará propiedades físicas como velocidad, movimiento, spin y release. Location+ estimará calidad de ubicación condicionada al tipo de lanzamiento, conteo y lado del bateador. Pitching+ combinará física, ubicación, arsenal y ambiente. Los resultados observados se utilizarán como objetivos de aprendizaje, nunca como entradas para predecir su propio resultado.')
p('Comenzaremos con un baseline por conteo y un modelo tabular de gradient boosting. Añadiremos arsenal y ambiente por etapas para medir su contribución. Las probabilidades distinguirán eventos por lanzamiento de tasas condicionadas, como whiff dado swing y chase dado lanzamiento fuera de zona.')
h('Ventajas que buscamos demostrar')
p('No conocemos los proyectos de los otros equipos. Los siguientes son diferenciadores propuestos, cuya ventaja deberá sostenerse con resultados y una demostración utilizable.')
table(['Diferenciador','Cómo lo demostraremos'],[
('Adaptación del repertorio','Comparar perfiles físicos y diferenciales respecto a la recta de referencia, además del score general del pitcher.'),
('Evidencia dentro del pitcher','Analizar el mismo pitcher, temporada y tipo de pitch entre ambientes, y después ajustar otras variables relevantes.'),
('Generalización verificable','Reportar desempeño en temporadas y pitchers no vistos, con particiones fijas y comparación contra baseline.'),
('Incertidumbre visible','Mostrar muestra disponible y rangos de incertidumbre; señalar cuándo el sistema no sostiene una recomendación.'),
('Decisión de scouting explicable','Presentar tres casos con comparación de candidatos, factores que explican el score y condiciones que cambiarían la recomendación.')],[2.15,4.75])
h('Validación y criterio de éxito')
p('Proponemos entrenar con 2024, seleccionar y calibrar con 2025, y reservar 2026 para evaluación final. Existen 249 IDs de pitcher en 2026 ausentes en los dos años previos, con 42,162 lanzamientos; servirán como prueba adicional si se confirma la estabilidad de los IDs.')
p('Mediremos calibración y Brier o log loss para probabilidades. Compararemos baseline, física, arsenal y ambiente sobre las mismas particiones. El criterio es mejorar frente al baseline y mantener utilidad fuera de muestra; no fijamos una mejora numérica sin experimentos. Park holdout requerirá identificar estadios.')

page('Organización del equipo y ejecución')
p('Los roles son una propuesta para asignar entre los cuatro integrantes según sus fortalezas. Cada responsable entrega resultados revisables y escribe la sección técnica de su área. La integración comienza desde el primer día.')
table(['Rol','Responsabilidad','Entrega al día 3'],[
('1 Datos y evaluación','Auditar etiquetas, preparar datos y fijar particiones. Revisar fugas de información.','Tabla preparada, reglas de limpieza y splits reproducibles.'),
('2 Modelado','Construir baseline y las tres capas. Evaluar y calibrar probabilidades.','Primer modelo comparado contra baseline.'),
('3 Arsenal y ambiente','Crear variables relacionales, estudiar diferencias ambientales y estimar incertidumbre.','Comparaciones por perfil y diseño de variables de arsenal.'),
('4 Producto y comunicación','Construir dashboard, integrar predicciones y coordinar documento y demo.','Prototipo con predicciones de ejemplo claramente identificadas.')],[1.45,3.15,2.3])
p('Asignación sugerida: la persona con mayor experiencia en ML toma modelado; quien tenga mayor atención al detalle y manejo de datos toma evaluación; quien prefiera estadística, física o béisbol toma ambiente; y quien domine visualización y comunicación toma producto. Los nombres quedan pendientes de acuerdo del equipo.')
table(['Periodo','Hito compartido'],[('Días 1 a 3','Resolver etiquetas prioritarias, definir targets y splits; baseline y prototipo.'),('Días 4 a 6','Construir las tres capas e incorporar arsenal sin información de prueba.'),('Días 7 a 9','Evaluar ambiente, calibración, generalización e incertidumbre.'),('Días 10 a 11','Integrar dashboard y seleccionar tres casos de scouting.'),('Día 12','Cerrar documento técnico y resultados reproducibles.'),('Días 13 y 14','Ensayar la entrevista, corregir fallos y preparar respaldo local.')],[1.45,5.45])
h('Acuerdos de trabajo')
p('Mantener una reunión diaria breve y un contrato común de columnas de entrada y salida. Datos revisa las particiones de modelado; modelado revisa las estimaciones ambientales; ambiente revisa las interpretaciones del dashboard. Si al día 4 no existe una evaluación completa, se reduce complejidad. Al día 10 se congela el alcance de modelos.')
p('El plan usa 14 días efectivos. La convocatoria compartida indica cierre de registro el 12 de octubre y entrega el 22 de octubre de 2026; el equipo debe confirmar su fecha operativa.')

page('Limitaciones y acciones prioritarias')
table(['Limitación verificada','Implicación y respuesta'],[
('Sin estadio ni altitud numérica','No identificar Harp Helú ni atribuir diferencias solo a altitud. Solicitar mapeo de partido a estadio, indicador Harp y metros de altitud.'),
('Sin fecha, orden de pitch ni ID de aparición al plato','No asumir secuencia cronológica ni calcular pitch anterior. Solicitar estos campos y definir el historial de arsenal disponible en cada evaluación.'),
('Sin ocupación de bases ni run value oficial','No reconstruir directamente el valor de carreras de 24 estados. Solicitar datos o acordar un proxy explícito; RunsScored no lo sustituye.'),
('Etiquetas contradictorias','is_hit_by_pitch siempre es cero pese a 2,280 HitByPitch en PitchCall. Hay 450 swings fallidos marcados como contacto. Acordar fuente canónica y reglas.'),
('Faltantes y valores sospechosos','3,272 filas sin pitcher; 3,020 sin altitud; 99 extensiones negativas; 6,669 alturas en plato iguales a cero. Revisar por campo y documentar exclusiones.'),
('Contacto y categorías escasas','ExitSpeed falta en 6.70% de las bolas en juego. Hay tipos con una sola fila. Evitar imputar cero y no hacer recomendaciones con muestras insuficientes.')],[2.25,4.65])
h('Qué se puede entregar aunque falten datos adicionales')
p('Un modelo físico, Location+ y un modelo combinado, comparaciones por categoría de altitud, evaluación por temporada y pitcher, y un dashboard que explique sus límites. El simulador específico del Harp, la validación por estadio y el cálculo estándar de valor de carreras dependen de información adicional. Una simulación ambiental será un escenario condicionado a supuestos, no un contrafactual causal demostrado.')
h('Primeras decisiones del equipo')
p('Asignar los cuatro responsables; enviar las preguntas de datos a los organizadores; fijar reglas para las etiquetas; definir una primera variable objetivo y congelar las particiones antes de experimentar. La prioridad es tener un flujo completo que vaya de los datos a una decisión explicable.')
h('Base de la propuesta')
p('Convocatoria y rúbrica proporcionadas por el equipo. Dataset Diablos - Hackathon 2026 (1).zip, tabla stuff_model_df.parquet y diccionario stuff_plus_data_dictionary.csv. Auditoría completa del 27 de septiembre de 2026 y comparación descriptiva por pitcher, temporada y tipo de lanzamiento. Los hallazgos son exploratorios y no constituyen validación predictiva.')
d.core_properties.title='Propuesta de evaluación de pitcheo para Diablos Rojos'
d.core_properties.subject='Proyecto de hackathon y plan de trabajo para cuatro integrantes'
d.core_properties.author='Equipo del hackathon'
d.save(out/'Propuesta_proyecto_Diablos.docx')
print(out/'Propuesta_proyecto_Diablos.docx')
